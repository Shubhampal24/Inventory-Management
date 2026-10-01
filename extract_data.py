import openpyxl
import json
from datetime import datetime
import math

wb = openpyxl.load_workbook(r'c:\Users\DELL\Downloads\Warehouse Rack Inventory.xlsx', data_only=True)

# 1. Racks
racks = []
ws_rack = wb['Rack Master']
for row in ws_rack.iter_rows(min_row=2, values_only=True):
    if not row[0]: continue
    rid = str(row[0]).strip()
    if rid in ['Level Code', 'GL1', 'GL2', 'GL3', 'SL1', 'SL2', 'SL3', 'SL4', 'SL5', 'SL6']: continue
    bayCount = 1
    if row[3] is not None:
        try: bayCount = int(float(row[3]))
        except: pass
    levels = []
    if row[4]: levels = [s.strip() for s in str(row[4]).split(',')]
    racks.append({
        'id': rid,
        'type': row[1] or 'Pallet Rack',
        'description': row[2] or '',
        'bayCount': bayCount,
        'levels': levels,
        'slots': [],
        'side': row[5] or 'Single-sided',
        'status': row[6] or 'Active'
    })
with open('src/lib/data/racks.js', 'w', encoding='utf-8') as f:
    f.write('export const racks = ' + json.dumps(racks, indent=2) + ';\n')

# 2. Materials
materials = []
ws_mat = wb['Material Master']
for row in ws_mat.iter_rows(min_row=2, values_only=True):
    if not row[0]: continue
    materials.append({
        'id': str(row[0]).strip(),
        'name': str(row[1]).strip() if row[1] else '',
        'category': str(row[2]).strip() if row[2] else '',
        'unit': str(row[3]).strip() if row[3] else '',
        'reorderLevel': float(row[4]) if row[4] is not None else 0
    })
with open('src/lib/data/materials.js', 'w', encoding='utf-8') as f:
    f.write('export const materials = ' + json.dumps(materials, indent=2) + ';\n')

# 3. Locations
locations = []
ws_loc = wb['Location Master']
for row in ws_loc.iter_rows(min_row=2, values_only=True):
    if not row[0]: continue
    locations.append({
        'rack': str(row[0]).strip(),
        'bay': int(row[1]) if row[1] is not None else 1,
        'level': str(row[2]).strip() if row[2] else '',
        'slot': str(row[3]).strip() if row[3] else '',
        'locationId': str(row[4]).strip() if row[4] else '',
        'materialId': str(row[6]).strip() if row[6] else '',
        'batch': '',
        'notes': ''
    })
with open('src/lib/data/locations.js', 'w', encoding='utf-8') as f:
    f.write('export const locations = ' + json.dumps(locations, indent=2) + ';\n')

# 4. Movements
movements = []
ws_mov = wb['Stock Movement']
for i, row in enumerate(ws_mov.iter_rows(min_row=2, values_only=True)):
    if not row[1] and not row[3]: continue
    dt = row[0]
    if isinstance(dt, datetime):
        dt_str = dt.strftime('%Y-%m-%d')
    else:
        dt_str = str(dt).split(' ')[0] if dt else datetime.now().strftime('%Y-%m-%d')
        
    quantity = 0
    if row[5] is not None:
        try: quantity = float(row[5])
        except ValueError: pass

    movements.append({
        'id': f'MOV_{i}',
        'date': dt_str,
        'materialId': str(row[1]).strip() if row[1] else '',
        'materialDesc': str(row[2]).strip() if row[2] else '',
        'locationId': str(row[3]).strip() if row[3] else '',
        'type': str(row[4]).strip() if row[4] else 'IN',
        'quantity': quantity,
        'unit': str(row[6]).strip() if row[6] else '',
        'reference': str(row[7]).strip() if row[7] else '',
        'user': str(row[8]).strip() if row[8] else '',
        'notes': str(row[9]).strip() if row[9] else ''
    })
with open('src/lib/data/movements.js', 'w', encoding='utf-8') as f:
    f.write('export const movements = ' + json.dumps(movements, indent=2) + ';\n')
